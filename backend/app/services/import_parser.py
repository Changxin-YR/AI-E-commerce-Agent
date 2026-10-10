"""Bounded text-only input. Never evaluate workbook or CSV expressions."""

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import PurePath
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring  # type: ignore[import-untyped]
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from app.core.errors import BusinessError
from app.schemas.imports import ParsedRow

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 64
MAX_CELL = 2000


def invalid_file(message: str) -> BusinessError:
    return BusinessError("invalid_file", message, 422)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, date | datetime):
        return value.isoformat()
    return str(value)


def check_archive(data: bytes) -> None:
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 256 or sum(entry.file_size for entry in entries) > 16 * 1024 * 1024:
            raise invalid_file("Excel 解压内容过大，请拆分文件")
        for entry in entries:
            name = entry.filename.lower()
            if any(
                part in name for part in ("vbaproject", "externallinks/", "embeddings/", "activex/")
            ):
                raise invalid_file("文件包含宏、外部链接或嵌入对象，请使用纯数据 .xlsx")
            if entry.flag_bits & 1:
                raise invalid_file("不支持加密 Excel 文件")
            if name.endswith(".rels"):
                root = fromstring(archive.read(entry))
                if any(node.attrib.get("TargetMode", "").lower() == "external" for node in root):
                    raise invalid_file("Excel 包含外部链接，请移除后再导入")
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                root = fromstring(archive.read(entry))
                for cell in root.iter(
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"
                ):
                    column, row_number = coordinate_from_string(cell.attrib["r"])
                    if column_index_from_string(column) > MAX_COLUMNS or row_number > MAX_ROWS + 1:
                        raise invalid_file("Excel 最多 64 列、2000 个数据行，请移除多余区域")


def parse_file(filename: str, data: bytes) -> tuple[list[str], list[ParsedRow], str]:
    if not data or len(data) > MAX_BYTES:
        raise invalid_file("请选择非空且不超过 2 MiB 的文件")
    extension = PurePath(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise invalid_file("仅支持 UTF-8 CSV 和无宏 .xlsx 文件")
    try:
        matrix: list[tuple[int, list[str], list[int]]]
        if extension == ".csv":
            source = data.decode("utf-8-sig")
            if "\x00" in source:
                raise invalid_file("CSV 含有无效控制字符")
            reader = csv.reader(StringIO(source, newline=""), strict=True)
            matrix = []
            start_line = 1
            for values in reader:
                matrix.append((start_line, values, []))
                start_line = reader.line_num + 1
                if len(matrix) > MAX_ROWS + 1:
                    raise invalid_file("每个文件最多 2000 个数据行，请拆分后导入")
            sheet_name = "CSV"
        else:
            check_archive(data)
            workbook = load_workbook(
                BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.worksheets) != 1:
                    raise invalid_file(
                        "请将要导入的工作表另存为单工作表 .xlsx；第一行保留唯一表头，"
                        "移除说明行及重复表头后重试"
                    )
                sheet = workbook.worksheets[0]
                # Ignore untrusted sheet dimensions, but cap actual iteration.
                sheet.reset_dimensions()
                matrix = []
                for number, cells in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS + 1), 1):
                    values = [cell_text(cell.value) for cell in cells]
                    while values and not values[-1]:
                        values.pop()
                    unsafe = [i for i, cell in enumerate(cells) if cell.data_type in {"f", "e"}]
                    matrix.append((number, values, unsafe))
                    if number > MAX_ROWS + 1:
                        raise invalid_file("每个文件最多 2000 个数据行，请拆分后导入")
                sheet_name = sheet.title
            finally:
                workbook.close()
    except BusinessError:
        raise
    except UnicodeDecodeError as error:
        raise invalid_file(
            "CSV 需要 UTF-8 编码（可带 BOM）；GB18030/GBK 文件请在表格软件中"
            "选择原编码打开，核对中文后另存为 CSV UTF-8 或单工作表 .xlsx"
        ) from error
    except Exception as error:
        # Parser errors must never echo source cells or internal paths.
        raise invalid_file("无法安全解析文件；请检查编码、格式或重新导出纯数据文件") from error
    if len(matrix) < 2:
        raise invalid_file("文件须有表头和至少一个数据行")
    headers = [value.strip() for value in matrix[0][1]]
    if not headers or len(headers) > MAX_COLUMNS or any(not h or len(h) > 120 for h in headers):
        raise invalid_file("表头须非空，最多 64 列，每个列名最多 120 字符")
    if len(set(headers)) != len(headers) or matrix[0][2]:
        raise invalid_file("表头重复或含公式，请修正列名")
    rows = []
    for number, values, unsafe in matrix[1:]:
        if len(values) > len(headers) or any(len(value) > MAX_CELL for value in values):
            raise invalid_file(f"源行 {number} 列数超出表头或单元格超过 2000 字符")
        rows.append(
            ParsedRow(
                row_number=number,
                values=values + [""] * (len(headers) - len(values)),
                unsafe_columns=unsafe,
            )
        )
    return headers, rows, sheet_name
